"""Import the user's ranked film and character list without inventing cast members."""
import json
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
source=ROOT/'视频排序(4).xlsx'
worksheet=openpyxl.load_workbook(source,data_only=True).active
identifiers=[
 ('dont-embarrass-us',['chovy-son','face-keeper']),
 ('trend-chasing-boss',['villain-calmer','trend-boss']),
 ('new-year-peace',['apocalypse-caterer']),
 ('not-a-coward',['dragon-office-worker','calm-responder']),
 ('never-give-up',['never-give-up']),
 ('too-much-thrill',['lucky-disaster']),
 ('royal-doorframe',['rule-breaker','social-shield']),
 ('royal-kitchen',['royal-client','royal-chef']),
 ('follow-your-heart',['plot-rewriter']),
 ('only-the-emperor',['hidden-player','flexible-challenger']),
 ('chaos-comedy',['chaos-director','chaos-witness']),
 ('public-roast',['public-roaster','quiet-bystander']),
 ('social-connection',['spotlight-extra']),
 ('mind-reading',['mind-reader','frustrated-courtier']),
 ('cosmic-stage',['stage-lead','cosmic-director']),
 ('reborn-revenge',['reborn-queen','queen-companion']),
 ('cultivation-support',['immortal-support']),
]
films=[]
for offset,(film_id,cast_ids) in enumerate(identifiers,2):
 title=worksheet.cell(offset,1).value
 count=worksheet.cell(offset,7).value
 names=[worksheet.cell(offset,col).value for col in [8,9] if worksheet.cell(offset,col).value]
 if count!=len(names) or count!=len(cast_ids):raise ValueError(f'Cast count mismatch at row {offset}: {title}')
 films.append({'id':film_id,'title':title,'sourceRow':offset,'rankTotal':worksheet.cell(offset,6).value,
               'characterCount':count,'cast':[{'id':id,'name':name} for id,name in zip(cast_ids,names)],
               'notes':worksheet.cell(offset,11).value or ''})
assert len(films)==17 and sum(f['characterCount'] for f in films)==28
(ROOT/'content/selected-films.json').write_text(json.dumps({'source':source.name,'films':films},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('Imported 17 selected films and 28 exactly named characters from the spreadsheet.')
