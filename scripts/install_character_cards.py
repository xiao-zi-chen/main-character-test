"""Install reviewed, individually generated role cards without changing videos."""
import datetime
import hashlib
import json
from pathlib import Path
import shutil

import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parent
GENERATED = PROJECT / 'output/character-cards-2026-10-27'
MEDIA = ROOT / 'public/media'


def sha(data):
    return hashlib.sha256(data).hexdigest()


def main():
    films = json.loads((ROOT / 'content/selected-films.json').read_text(encoding='utf-8'))['films']
    manifest_path = MEDIA / 'manifest.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    workbook = openpyxl.load_workbook(ROOT / '视频排序(4).xlsx', read_only=True, data_only=True)
    sheet = workbook.active
    cast = []
    for film in films:
        row = film['sourceRow']
        assert sheet.cell(row, 1).value == film['title']
        names = [sheet.cell(row, column).value for column in (8, 9) if sheet.cell(row, column).value]
        assert names == [person['name'] for person in film['cast']]
        cast.extend(film['cast'])
    workbook.close()
    assert len(cast) == 28 and {p['id'] for p in cast} == set(manifest)
    assert not GENERATED.is_symlink()
    reviewed = json.loads((GENERATED / 'review.json').read_text(encoding='utf-8'))
    assert set(reviewed) == set(manifest), 'Review must cover all 28 characters'
    pending = []
    for person in cast:
        key = person['id']
        path = GENERATED / f'{key}.png'
        assert path.is_file() and not path.is_symlink()
        data = path.read_bytes()
        digest = sha(data)
        review = reviewed[key]
        assert review['title'] == person['name'] and review['sha256'] == digest
        assert review['singleCharacter'] and review['titleChecked'] and review['identityChecked']
        with Image.open(path) as image:
            image.load()
            assert image.format == 'PNG' and image.width >= 900 and image.height > image.width
            assert abs(image.width / image.height - 2 / 3) < .01, key
            size = [image.width, image.height]
        filename = f'{key}-card-{digest[:12]}.png'
        pending.append((key, path, filename, digest, size, person['name']))

    backup = PROJECT / '.local/character-cards-oct27' / ('before-install-' + datetime.datetime.now(datetime.timezone.utc).strftime('%Y%m%dT%H%M%S%fZ'))
    backup.mkdir(parents=True)
    shutil.copy2(manifest_path, backup / 'manifest.json')
    for key, _, _, _, _, _ in pending:
        old = MEDIA / manifest[key]['card']
        assert old.is_file() and old.parent == MEDIA and not old.is_symlink()
        shutil.copy2(old, backup / old.name)
    records = []
    for key, source, filename, digest, size, title in pending:
        target = MEDIA / filename
        if target.exists():
            assert sha(target.read_bytes()) == digest
        else:
            shutil.copy2(source, target)
        entry = manifest[key]
        entry.update(card=filename, cardSha256=digest, cardSize=size, cardTitle=title,
                     cardSource='imagegen-film-reference')
        records.append({'id': key, 'title': title, 'file': filename, 'sha256': digest, 'size': size})
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    report = {'cards': records, 'backup': str(backup), 'videosChanged': False, 'cloudChanged': False}
    (backup.parent / 'install-result.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Installed {len(records)} reviewed single-character cards. Backup: {backup}')


if __name__ == '__main__':
    main()
